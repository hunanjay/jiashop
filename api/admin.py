import csv
import io
import os
from datetime import datetime, timedelta

from flask import Blueprint, jsonify, make_response, request
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage
from sqlalchemy import func
from werkzeug.security import generate_password_hash

from api.dependencies import require_permission
from db.models import CommissionRecord, Customer, Shop, User, Visit, db
from db.models import Order
from services.admin_service import get_role_by_name, list_roles, update_user_role
from services.oss_service import download_image_bytes, get_signed_url, upload_base64_to_oss
from api.dependencies import get_current_user


admin_bp = Blueprint("admin", __name__, url_prefix="/api/admin")
workspace_bp = Blueprint("workspace", __name__, url_prefix="/api/workspace")


def _serialize_user(user: User):
    return {
        "id": user.id,
        "username": user.username,
        "email": user.email,
        "phone": user.phone,
        "role": user.role.name if user.role else None,
        "created_at": user.created_at.isoformat() if user.created_at else None,
    }


def _serialize_customer(customer: Customer):
    return {
        "id": customer.id,
        "company_name": customer.company_name,
        "purchaser": customer.purchaser,
        "phone": customer.phone,
        "shipping_address": customer.shipping_address,
        "owner_id": customer.owner_id,
        "created_at": customer.created_at.isoformat() if customer.created_at else None,
        "updated_at": customer.updated_at.isoformat() if customer.updated_at else None,
    }


@admin_bp.route("/users", methods=["GET"])
@require_permission("/api/admin/users", "GET")
def list_users():
    """
    List users
    ---
    tags:
      - Admin
    security:
      - bearerAuth: []
    responses:
      200:
        description: User list
      403:
        description: Forbidden
    """
    users = User.query.order_by(User.created_at.desc()).all()
    return jsonify([_serialize_user(user) for user in users])


@admin_bp.route("/users", methods=["POST"])
@require_permission("/api/admin/users", "POST")
def create_user():
    data = request.get_json(silent=True) or {}
    username = (data.get("username") or "").strip()
    email = (data.get("email") or "").strip()
    phone = (data.get("phone") or "").strip() or None
    password = (data.get("password") or "").strip()
    role_name = (data.get("role") or "user").strip()

    if not username or not email or not password:
        return jsonify({"error": "username, email and password are required"}), 400

    if User.query.filter((User.username == username) | (User.email == email)).first():
        return jsonify({"error": "username or email already exists"}), 400
    if phone and User.query.filter(User.phone == phone).first():
        return jsonify({"error": "phone already exists"}), 400

    role = get_role_by_name(role_name)
    if not role:
        return jsonify({"error": "Role not found", "role": role_name}), 400

    user = User(
        username=username,
        email=email,
        phone=phone,
        role_id=role.id,
        password_hash=generate_password_hash(password),
    )
    db.session.add(user)
    db.session.commit()
    return jsonify(_serialize_user(user)), 201


@admin_bp.route("/roles", methods=["GET"])
@require_permission("/api/admin/roles", "GET")
def get_roles():
    """
    List roles
    ---
    tags:
      - Admin
    security:
      - bearerAuth: []
    responses:
      200:
        description: Role list
      403:
        description: Forbidden
    """
    roles = list_roles()
    return jsonify([{"id": role.id, "name": role.name} for role in roles])


@admin_bp.route("/users/<user_id>/role", methods=["PATCH"])
@require_permission("/api/admin/users/:user_id/role", "PATCH")
def change_user_role(user_id):
    """
    Update user role
    ---
    tags:
      - Admin
    security:
      - bearerAuth: []
    parameters:
      - in: path
        name: user_id
        required: true
        type: string
      - in: body
        name: body
        required: true
        schema:
          type: object
          required:
            - role
          properties:
            role:
              type: string
    responses:
      200:
        description: Role updated
      400:
        description: Invalid role
      404:
        description: User not found
      403:
        description: Forbidden
    """
    data = request.get_json(silent=True) or {}
    role_name = (data.get("role") or "").strip()

    if not role_name:
        return jsonify({"error": "role is required"}), 400

    role = get_role_by_name(role_name)
    if not role:
        return jsonify({"error": "Role not found", "role": role_name}), 400

    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "User not found"}), 404

    actor = get_current_user()
    updated_user = update_user_role(user, role, actor.id if actor else None)
    return jsonify(_serialize_user(updated_user))


@admin_bp.route("/users/<user_id>", methods=["PATCH"])
@require_permission("/api/admin/users/:user_id", "PATCH")
def update_user(user_id):
    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "User not found"}), 404

    data = request.get_json(silent=True) or {}
    if "username" in data:
        username = (data.get("username") or "").strip()
        if not username:
            return jsonify({"error": "username is required"}), 400
        user.username = username
    if "email" in data:
        email = (data.get("email") or "").strip()
        if not email:
            return jsonify({"error": "email is required"}), 400
        user.email = email
    if "phone" in data:
        phone = (data.get("phone") or "").strip() or None
        if phone and User.query.filter(User.phone == phone, User.id != user.id).first():
            return jsonify({"error": "phone already exists"}), 400
        user.phone = phone
    if "role" in data:
        role = get_role_by_name((data.get("role") or "").strip())
        if not role:
            return jsonify({"error": "Role not found"}), 400
        user.role_id = role.id

    db.session.add(user)
    db.session.commit()
    return jsonify(_serialize_user(user))


@admin_bp.route("/users/<user_id>/password", methods=["PATCH"])
@require_permission("/api/admin/users/:user_id/password", "PATCH")
def reset_user_password(user_id):
    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "User not found"}), 404

    data = request.get_json(silent=True) or {}
    password = (data.get("password") or "").strip()
    if not password:
        return jsonify({"error": "password is required"}), 400

    user.set_password(password)
    db.session.add(user)
    db.session.commit()
    return jsonify(_serialize_user(user))


@admin_bp.route("/users/<user_id>", methods=["DELETE"])
@require_permission("/api/admin/users/:user_id", "DELETE")
def delete_user(user_id):
    user = User.query.get(user_id)
    if not user:
        return jsonify({"error": "User not found"}), 404

    db.session.delete(user)
    db.session.commit()
    return jsonify({"deleted": True, "id": user_id})


@admin_bp.route("/customers", methods=["GET"])
@require_permission("/api/admin/customers", "GET")
def list_customers():
    customers = Customer.query.order_by(Customer.created_at.desc()).all()
    return jsonify([_serialize_customer(customer) for customer in customers])


@admin_bp.route("/customers", methods=["POST"])
@require_permission("/api/admin/customers", "POST")
def create_customer():
    data = request.get_json(silent=True) or {}
    company_name = (data.get("company_name") or "").strip()
    if not company_name:
        return jsonify({"error": "company_name is required"}), 400

    customer = Customer(
        company_name=company_name,
        purchaser=(data.get("purchaser") or "").strip() or None,
        phone=(data.get("phone") or "").strip() or None,
        shipping_address=(data.get("shipping_address") or "").strip() or None,
        owner_id=(data.get("owner_id") or "").strip() or None,
    )
    db.session.add(customer)
    db.session.commit()
    return jsonify(_serialize_customer(customer)), 201


@admin_bp.route("/customers/<customer_id>", methods=["PATCH"])
@require_permission("/api/admin/customers/:customer_id", "PATCH")
def update_customer(customer_id):
    customer = Customer.query.get(customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    data = request.get_json(silent=True) or {}
    if "company_name" in data:
        company_name = (data.get("company_name") or "").strip()
        if not company_name:
            return jsonify({"error": "company_name is required"}), 400
        customer.company_name = company_name
    if "purchaser" in data:
        customer.purchaser = (data.get("purchaser") or "").strip() or None
    if "phone" in data:
        customer.phone = (data.get("phone") or "").strip() or None
    if "shipping_address" in data:
        customer.shipping_address = (data.get("shipping_address") or "").strip() or None
    if "owner_id" in data:
        customer.owner_id = (data.get("owner_id") or "").strip() or None

    db.session.add(customer)
    db.session.commit()
    return jsonify(_serialize_customer(customer))


@admin_bp.route("/customers/<customer_id>", methods=["DELETE"])
@require_permission("/api/admin/customers/:customer_id", "DELETE")
def delete_customer(customer_id):
    customer = Customer.query.get(customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    db.session.delete(customer)
    db.session.commit()
    return jsonify({"deleted": True, "id": customer_id})


@workspace_bp.route("/customers", methods=["GET"])
@require_permission("/api/workspace/customers", "GET")
def workspace_list_customers():
    current_user = get_current_user()
    query = Customer.query
    if current_user and current_user.role and current_user.role.name and current_user.role.name.lower() == "user":
        query = query.filter(Customer.owner_id == current_user.id)
    customers = query.order_by(Customer.created_at.desc()).all()
    return jsonify([_serialize_customer(customer) for customer in customers])


@workspace_bp.route("/customers", methods=["POST"])
@require_permission("/api/workspace/customers", "POST")
def workspace_create_customer():
    data = request.get_json(silent=True) or {}
    company_name = (data.get("company_name") or "").strip()
    if not company_name:
        return jsonify({"error": "company_name is required"}), 400

    current_user = get_current_user()
    customer = Customer(
        company_name=company_name,
        purchaser=(data.get("purchaser") or "").strip() or None,
        phone=(data.get("phone") or "").strip() or None,
        shipping_address=(data.get("shipping_address") or "").strip() or None,
        owner_id=current_user.id if current_user else None,
    )
    db.session.add(customer)
    db.session.commit()
    return jsonify(_serialize_customer(customer)), 201


@workspace_bp.route("/customers/<customer_id>", methods=["PATCH"])
@require_permission("/api/workspace/customers/:customer_id", "PATCH")
def workspace_update_customer(customer_id):
    customer = Customer.query.get(customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    current_user = get_current_user()
    if current_user and customer.owner_id and customer.owner_id != current_user.id:
        return jsonify({"error": "Insufficient permissions"}), 403

    data = request.get_json(silent=True) or {}
    if "company_name" in data:
        company_name = (data.get("company_name") or "").strip()
        if not company_name:
            return jsonify({"error": "company_name is required"}), 400
        customer.company_name = company_name
    if "purchaser" in data:
        customer.purchaser = (data.get("purchaser") or "").strip() or None
    if "phone" in data:
        customer.phone = (data.get("phone") or "").strip() or None
    if "shipping_address" in data:
        customer.shipping_address = (data.get("shipping_address") or "").strip() or None

    db.session.add(customer)
    db.session.commit()
    return jsonify(_serialize_customer(customer))


@workspace_bp.route("/customers/<customer_id>", methods=["DELETE"])
@require_permission("/api/workspace/customers/:customer_id", "DELETE")
def workspace_delete_customer(customer_id):
    customer = Customer.query.get(customer_id)
    if not customer:
        return jsonify({"error": "Customer not found"}), 404

    current_user = get_current_user()
    if current_user and customer.owner_id and customer.owner_id != current_user.id:
        return jsonify({"error": "Insufficient permissions"}), 403

    db.session.delete(customer)
    db.session.commit()
    return jsonify({"deleted": True, "id": customer_id})


COMMISSION_PLATFORMS = ("淘宝", "小红书")


def _remember_shop(owner_id, shop_name):
    """Adds shop_name to the owner's known-shops list if it's new, so it can be suggested next time."""
    if not owner_id or not shop_name:
        return
    exists = Shop.query.filter_by(owner_id=owner_id, name=shop_name).first()
    if not exists:
        db.session.add(Shop(name=shop_name, owner_id=owner_id))
        db.session.commit()


def _parse_order_time(value):
    if not value:
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return None


def _serialize_commission(record: CommissionRecord):
    return {
        "id": record.id,
        "platform": record.platform,
        "shop_name": record.shop_name,
        "product_name": record.product_name,
        "product_image": get_signed_url(record.product_image),
        "order_no": record.order_no,
        "order_time": record.order_time.isoformat() if record.order_time else None,
        "quantity": record.quantity,
        "amount": record.amount,
        "commission": record.commission,
        "owner_id": record.owner_id,
        "created_at": record.created_at.isoformat() if record.created_at else None,
        "updated_at": record.updated_at.isoformat() if record.updated_at else None,
    }


def _extract_commission_payload(data, current_user):
    platform = (data.get("platform") or "").strip() or COMMISSION_PLATFORMS[0]
    shop_name = (data.get("shop_name") or "").strip()
    product_name = (data.get("product_name") or "").strip()
    order_no = (data.get("order_no") or "").strip()
    if platform not in COMMISSION_PLATFORMS:
        return None, jsonify({"error": f"platform must be one of {COMMISSION_PLATFORMS}"}), 400
    if not shop_name or not product_name or not order_no:
        return None, jsonify({"error": "shop_name, product_name and order_no are required"}), 400

    try:
        quantity = int(data.get("quantity", 1))
        amount = float(data.get("amount", 0))
        commission = float(data.get("commission", 0))
    except (TypeError, ValueError):
        return None, jsonify({"error": "quantity, amount and commission must be numbers"}), 400

    payload = {
        "platform": platform,
        "shop_name": shop_name,
        "product_name": product_name,
        "order_no": order_no,
        "order_time": _parse_order_time(data.get("order_time")),
        "quantity": quantity,
        "amount": amount,
        "commission": commission,
    }
    if "product_image" in data:
        payload["product_image"] = upload_base64_to_oss(data.get("product_image"), user_id=current_user.id if current_user else "system") or None
    return payload, None, None


def _filter_commissions(query):
    """Applies the shared ?platform=&month=YYYY-MM query params used by list and export."""
    platform = (request.args.get("platform") or "").strip()
    if platform and platform in COMMISSION_PLATFORMS:
        query = query.filter(CommissionRecord.platform == platform)

    month = (request.args.get("month") or "").strip()
    if month:
        try:
            start = datetime.strptime(month, "%Y-%m")
        except ValueError:
            return query
        end = (start + timedelta(days=32)).replace(day=1)
        query = query.filter(CommissionRecord.order_time >= start, CommissionRecord.order_time < end)
    return query


COMMISSION_IMAGE_DISPLAY_PX = 60  # on-screen size in the cell; the embedded file stays full-res


def _commission_xlimage(image_path):
    """Downloads a commission record's product image at full resolution and
    returns it as an openpyxl Image ready to anchor into a cell, or None if
    unavailable. The PNG bytes are kept at full resolution (not compressed) —
    only the on-screen display width/height is scaled down so the picture
    actually sits inside its row instead of overflowing across the sheet."""
    raw = download_image_bytes(image_path)
    if not raw:
        return None
    try:
        pil_img = PILImage.open(io.BytesIO(raw)).convert("RGB")
        buf = io.BytesIO()
        pil_img.save(buf, format="PNG")
        buf.seek(0)
        xl_img = XLImage(buf)
        scale = COMMISSION_IMAGE_DISPLAY_PX / max(pil_img.width, pil_img.height)
        xl_img.width = pil_img.width * scale
        xl_img.height = pil_img.height * scale
        return xl_img
    except Exception as e:
        print(f"Failed to embed image for {image_path}: {e}")
        return None


def _export_commissions_xlsx(records, include_owner, owner_map):
    header = ["平台", "店铺名", "产品图片", "产品名称", "订单号", "下单时间", "数量", "金额", "佣金"]
    if include_owner:
        header.append("归属人员")
    header.append("创建时间")
    image_col = header.index("产品图片") + 1  # 1-indexed for openpyxl

    wb = Workbook()
    ws = wb.active
    ws.title = "佣金记录"
    ws.append(header)
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.column_dimensions[get_column_letter(image_col)].width = COMMISSION_IMAGE_DISPLAY_PX / 7 + 2

    for row_idx, record in enumerate(records, start=2):
        row = [
            record.platform,
            record.shop_name,
            "",  # image goes in as a floating picture, not cell text
            record.product_name,
            record.order_no,
            record.order_time.isoformat(sep=" ") if record.order_time else "",
            record.quantity,
            record.amount,
            record.commission,
        ]
        if include_owner:
            row.append(owner_map.get(record.owner_id, record.owner_id or ""))
        row.append(record.created_at.isoformat(sep=" ") if record.created_at else "")
        ws.append(row)

        image = _commission_xlimage(record.product_image)
        if image:
            ws.row_dimensions[row_idx].height = COMMISSION_IMAGE_DISPLAY_PX * 0.75 + 5
            ws.add_image(image, f"{get_column_letter(image_col)}{row_idx}")

    for col_idx, title in enumerate(header, start=1):
        if col_idx == image_col:
            continue
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, len(title) + 4)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    response = make_response(output.getvalue())
    response.headers["Content-Type"] = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    response.headers["Content-Disposition"] = 'attachment; filename="commissions.xlsx"'
    return response


@admin_bp.route("/commissions", methods=["GET"])
@require_permission("/api/admin/commissions", "GET")
def list_commissions():
    query = _filter_commissions(CommissionRecord.query)
    records = query.order_by(CommissionRecord.created_at.desc()).all()
    return jsonify([_serialize_commission(record) for record in records])


@admin_bp.route("/commissions/export", methods=["GET"])
@require_permission("/api/admin/commissions/export", "GET")
def export_commissions():
    query = _filter_commissions(CommissionRecord.query)
    records = query.order_by(CommissionRecord.created_at.desc()).all()
    owner_map = {user.id: user.username for user in User.query.all()}
    return _export_commissions_xlsx(records, include_owner=True, owner_map=owner_map)


@admin_bp.route("/commissions", methods=["POST"])
@require_permission("/api/admin/commissions", "POST")
def create_commission():
    data = request.get_json(silent=True) or {}
    current_user = get_current_user()
    payload, error_response, status = _extract_commission_payload(data, current_user)
    if error_response:
        return error_response, status

    payload["owner_id"] = (data.get("owner_id") or "").strip() or (current_user.id if current_user else None)
    record = CommissionRecord(**payload)
    db.session.add(record)
    db.session.commit()
    _remember_shop(record.owner_id, record.shop_name)
    return jsonify(_serialize_commission(record)), 201


@admin_bp.route("/commissions/<record_id>", methods=["PATCH"])
@require_permission("/api/admin/commissions/:record_id", "PATCH")
def update_commission(record_id):
    record = CommissionRecord.query.get(record_id)
    if not record:
        return jsonify({"error": "Commission record not found"}), 404

    data = request.get_json(silent=True) or {}
    current_user = get_current_user()
    payload, error_response, status = _extract_commission_payload(data, current_user)
    if error_response:
        return error_response, status

    for key, value in payload.items():
        setattr(record, key, value)
    if "owner_id" in data:
        record.owner_id = (data.get("owner_id") or "").strip() or None

    db.session.add(record)
    db.session.commit()
    _remember_shop(record.owner_id, record.shop_name)
    return jsonify(_serialize_commission(record))


@admin_bp.route("/commissions/<record_id>", methods=["DELETE"])
@require_permission("/api/admin/commissions/:record_id", "DELETE")
def delete_commission(record_id):
    record = CommissionRecord.query.get(record_id)
    if not record:
        return jsonify({"error": "Commission record not found"}), 404

    db.session.delete(record)
    db.session.commit()
    return jsonify({"deleted": True, "id": record_id})


@workspace_bp.route("/commissions", methods=["GET"])
@require_permission("/api/workspace/commissions", "GET")
def workspace_list_commissions():
    current_user = get_current_user()
    query = CommissionRecord.query
    if current_user and current_user.role and current_user.role.name and current_user.role.name.lower() == "user":
        query = query.filter(CommissionRecord.owner_id == current_user.id)
    query = _filter_commissions(query)
    records = query.order_by(CommissionRecord.created_at.desc()).all()
    return jsonify([_serialize_commission(record) for record in records])


@workspace_bp.route("/commissions/export", methods=["GET"])
@require_permission("/api/workspace/commissions/export", "GET")
def workspace_export_commissions():
    current_user = get_current_user()
    query = CommissionRecord.query
    is_plain_user = current_user and current_user.role and current_user.role.name and current_user.role.name.lower() == "user"
    if is_plain_user:
        query = query.filter(CommissionRecord.owner_id == current_user.id)
    query = _filter_commissions(query)
    records = query.order_by(CommissionRecord.created_at.desc()).all()
    owner_map = {} if is_plain_user else {user.id: user.username for user in User.query.all()}
    return _export_commissions_xlsx(records, include_owner=not is_plain_user, owner_map=owner_map)


@workspace_bp.route("/commissions", methods=["POST"])
@require_permission("/api/workspace/commissions", "POST")
def workspace_create_commission():
    data = request.get_json(silent=True) or {}
    current_user = get_current_user()
    payload, error_response, status = _extract_commission_payload(data, current_user)
    if error_response:
        return error_response, status

    payload["owner_id"] = current_user.id if current_user else None
    record = CommissionRecord(**payload)
    db.session.add(record)
    db.session.commit()
    _remember_shop(record.owner_id, record.shop_name)
    return jsonify(_serialize_commission(record)), 201


@workspace_bp.route("/commissions/<record_id>", methods=["PATCH"])
@require_permission("/api/workspace/commissions/:record_id", "PATCH")
def workspace_update_commission(record_id):
    record = CommissionRecord.query.get(record_id)
    if not record:
        return jsonify({"error": "Commission record not found"}), 404

    current_user = get_current_user()
    if current_user and record.owner_id and record.owner_id != current_user.id:
        return jsonify({"error": "Insufficient permissions"}), 403

    data = request.get_json(silent=True) or {}
    payload, error_response, status = _extract_commission_payload(data, current_user)
    if error_response:
        return error_response, status

    for key, value in payload.items():
        setattr(record, key, value)

    db.session.add(record)
    db.session.commit()
    _remember_shop(record.owner_id, record.shop_name)
    return jsonify(_serialize_commission(record))


@workspace_bp.route("/commissions/<record_id>", methods=["DELETE"])
@require_permission("/api/workspace/commissions/:record_id", "DELETE")
def workspace_delete_commission(record_id):
    record = CommissionRecord.query.get(record_id)
    if not record:
        return jsonify({"error": "Commission record not found"}), 404

    current_user = get_current_user()
    if current_user and record.owner_id and record.owner_id != current_user.id:
        return jsonify({"error": "Insufficient permissions"}), 403

    db.session.delete(record)
    db.session.commit()
    return jsonify({"deleted": True, "id": record_id})


@workspace_bp.route("/shops", methods=["GET"])
@require_permission("/api/workspace/shops", "GET")
def workspace_list_shops():
    current_user = get_current_user()
    shops = Shop.query.filter_by(owner_id=current_user.id if current_user else None).order_by(Shop.name).all()
    return jsonify([shop.name for shop in shops])


@admin_bp.route("/stats", methods=["GET"])
@require_permission("/api/admin/stats", "GET")
def admin_stats():
    now = datetime.utcnow()
    week_start = now - timedelta(days=6)
    month_start = now - timedelta(days=29)

    status_rows = db.session.query(Order.status, func.count(Order.id)).group_by(Order.status).all()
    owner_rows = (
        db.session.query(User.username, func.coalesce(func.sum(Order.total_price), 0), func.count(Order.id))
        .outerjoin(Order, Order.owner_id == User.id)
        .group_by(User.username)
        .order_by(func.coalesce(func.sum(Order.total_price), 0).desc())
        .all()
    )
    customer_owner_rows = (
        db.session.query(User.username, func.count(Customer.id))
        .outerjoin(Customer, Customer.owner_id == User.id)
        .group_by(User.username)
        .order_by(func.count(Customer.id).desc())
        .all()
    )
    week_orders = Order.query.filter(Order.created_at >= week_start).all()
    month_orders = Order.query.filter(Order.created_at >= month_start).all()

    daily_trend = []
    for offset in range(7):
        day = (week_start + timedelta(days=offset)).date()
        day_orders = [order for order in week_orders if order.created_at and order.created_at.date() == day]
        daily_trend.append(
            {
                "label": day.strftime("%m-%d"),
                "orders": len(day_orders),
                "sales_total": float(sum(float(order.total_price or 0) for order in day_orders)),
            }
        )

    weekly_summary = []
    for week_index in range(4):
        start = month_start + timedelta(days=week_index * 7)
        end = start + timedelta(days=7)
        chunk_orders = [order for order in month_orders if order.created_at and start <= order.created_at < end]
        weekly_summary.append(
            {
                "label": f"W{week_index + 1}",
                "orders": len(chunk_orders),
                "sales_total": float(sum(float(order.total_price or 0) for order in chunk_orders)),
            }
        )

    # ponytail: 统计按 UTC+STATS_TZ_OFFSET 分天，默认东八区；部署在别的时区改环境变量即可
    tz_offset = timedelta(hours=int(os.environ.get("STATS_TZ_OFFSET", 8)))
    visit_rows = Visit.query.filter(Visit.created_at >= week_start).all()
    visit_trend = []
    for offset in range(7):
        day = (week_start + timedelta(days=offset) + tz_offset).date()
        day_visits = [visit for visit in visit_rows if visit.created_at and (visit.created_at + tz_offset).date() == day]
        visit_trend.append(
            {
                "label": day.strftime("%m-%d"),
                "pv": len(day_visits),
                "uv": len({visit.device_id for visit in day_visits}),
            }
        )
    today = (now + tz_offset).date()
    today_visits = [visit for visit in visit_rows if visit.created_at and (visit.created_at + tz_offset).date() == today]

    monthly_summary = {
        "orders": len(month_orders),
        "sales_total": float(sum(float(order.total_price or 0) for order in month_orders)),
    }
    return jsonify(
        {
            "users": User.query.count(),
            "customers": Customer.query.count(),
            "orders": Order.query.count(),
            "sales_total": db.session.query(func.coalesce(func.sum(Order.total_price), 0)).scalar() or 0,
            "pending_orders": Order.query.filter_by(status="Pending").count(),
            "processing_orders": Order.query.filter_by(status="Processing").count(),
            "completed_orders": Order.query.filter_by(status="Completed").count(),
            "status_distribution": [{"status": status, "count": count} for status, count in status_rows],
            "sales_ranking": [
                {"username": username, "sales_total": float(sales_total or 0), "order_count": int(order_count or 0)}
                for username, sales_total, order_count in owner_rows
            ],
            "customer_summary": {
                "recent": [
                    {
                        "id": customer.id,
                        "company_name": customer.company_name,
                        "purchaser": customer.purchaser,
                        "owner_id": customer.owner_id,
                        "created_at": customer.created_at.isoformat() if customer.created_at else None,
                    }
                    for customer in Customer.query.order_by(Customer.created_at.desc()).limit(5).all()
                ],
                "owner_distribution": [
                    {"username": username or "Unassigned", "count": int(count or 0)} for username, count in customer_owner_rows
                ],
            },
            "visits": {
                "total_pv": Visit.query.count(),
                "total_uv": db.session.query(func.count(func.distinct(Visit.device_id))).scalar() or 0,
                "today_pv": len(today_visits),
                "today_uv": len({visit.device_id for visit in today_visits}),
                "daily_trend": visit_trend,
            },
            "daily_trend": daily_trend,
            "weekly_summary": weekly_summary,
            "monthly_summary": monthly_summary,
        }
    )


@admin_bp.route("/export", methods=["GET"])
@require_permission("/api/admin/export", "GET")
def export_data():
    resource = (request.args.get("resource") or "orders").strip().lower()
    output = io.StringIO()
    writer = csv.writer(output)
    owner_map = {user.id: user.username for user in User.query.all()}

    if resource == "customers":
        writer.writerow(["id", "company_name", "purchaser", "phone", "shipping_address", "owner_username", "created_at"])
        for customer in Customer.query.order_by(Customer.created_at.desc()).all():
            writer.writerow([
                customer.id,
                customer.company_name,
                customer.purchaser or "",
                customer.phone or "",
                customer.shipping_address or "",
                owner_map.get(customer.owner_id, customer.owner_id or ""),
                customer.created_at.isoformat() if customer.created_at else "",
            ])
        filename = "customers.csv"
    elif resource == "users":
        writer.writerow(["id", "username", "email", "phone", "role", "created_at"])
        for user in User.query.order_by(User.created_at.desc()).all():
            writer.writerow([
                user.id,
                user.username,
                user.email,
                user.phone or "",
                user.role.name if user.role else "",
                user.created_at.isoformat() if user.created_at else "",
            ])
        filename = "users.csv"
    else:
        writer.writerow(["id", "customer_name", "status", "total_price", "customer_phone", "shipping_address", "owner_username", "created_at"])
        for order in Order.query.order_by(Order.created_at.desc()).all():
            writer.writerow([
                order.id,
                order.customer_name,
                order.status,
                order.total_price,
                order.customer_phone or "",
                order.shipping_address or "",
                owner_map.get(order.owner_id, order.owner_id or ""),
                order.created_at.isoformat() if order.created_at else "",
            ])
        filename = "orders.csv"

    response = make_response(output.getvalue())
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    response.headers["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
